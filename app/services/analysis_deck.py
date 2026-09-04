import os
from datetime import datetime

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from app.database.db import SessionLocal
from app.database.models import Bill, BillItem, Product


STORE_NAME = "NEBULA SUPERMARKET"

GENERATED_DIR = "generated"
CHART_DIR = os.path.join(GENERATED_DIR, "charts")


def format_money(amount):
    if amount is None:
        amount = 0.0

    return f"₹{float(amount):,.2f}"


def add_title(slide, title, subtitle=None):
    """
    Add a title to a slide.
    """

    title_box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.7)
    )

    title_frame = title_box.text_frame
    title_frame.clear()

    paragraph = title_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.6),
            Inches(1.0),
            Inches(12),
            Inches(0.4)
        )

        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.clear()

        paragraph = subtitle_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.size = Pt(12)


def add_text_box(slide, text, left, top, width, height, font_size=18):
    """
    Add a normal text box.
    """

    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(font_size)

    return box


def create_payment_chart(payment_data):
    """
    Create payment-mode chart.
    """

    os.makedirs(CHART_DIR, exist_ok=True)

    labels = []
    values = []

    for mode, amount in payment_data.items():
        labels.append(mode.upper())
        values.append(float(amount))

    chart_path = os.path.join(
        CHART_DIR,
        "payment_analysis.png"
    )

    plt.figure(figsize=(8, 5))

    if any(values):
        plt.bar(labels, values)

    plt.title("Sales by Payment Mode")
    plt.xlabel("Payment Mode")
    plt.ylabel("Sales Amount (INR)")
    plt.tight_layout()

    plt.savefig(chart_path, dpi=150)
    plt.close()

    return chart_path


def create_inventory_chart(inventory_data):
    """
    Create inventory stock chart.
    """

    os.makedirs(CHART_DIR, exist_ok=True)

    names = []
    quantities = []

    for item in inventory_data[:10]:
        names.append(item["name"])
        quantities.append(float(item["quantity"]))

    chart_path = os.path.join(
        CHART_DIR,
        "inventory_analysis.png"
    )

    plt.figure(figsize=(9, 5))

    if quantities:
        plt.bar(names, quantities)

    plt.title("Current Inventory - Top Products")
    plt.xlabel("Product")
    plt.ylabel("Stock Quantity")

    plt.xticks(
        rotation=35,
        ha="right"
    )

    plt.tight_layout()

    plt.savefig(chart_path, dpi=150)
    plt.close()

    return chart_path


def generate_business_insights(
    total_sales,
    total_bills,
    average_bill,
    credit_sales,
    low_stock_count,
    total_products
):
    """
    Generate simple business insights from database values.
    """

    insights = []

    if total_sales > 0:
        insights.append(
            f"Total finalized sales are {format_money(total_sales)}."
        )
    else:
        insights.append(
            "No finalized sales are available for the analysis period."
        )

    if total_bills > 0:
        insights.append(
            f"The supermarket processed {total_bills} finalized bill(s) "
            f"with an average bill value of {format_money(average_bill)}."
        )

    if credit_sales > 0:
        insights.append(
            f"Credit sales amount to {format_money(credit_sales)}, "
            "so khata balances should be monitored."
        )
    else:
        insights.append(
            "There are currently no credit sales in the analysis period."
        )

    if low_stock_count > 0:
        insights.append(
            f"{low_stock_count} product(s) are at or below their "
            "reorder level and should be reviewed."
        )
    else:
        insights.append(
            "No products are currently below their reorder threshold."
        )

    if total_products > 0:
        insights.append(
            f"The inventory currently contains {total_products} "
            "product SKU(s)."
        )

    return insights


def generate_analysis_deck():
    """
    Generate a PowerPoint business analysis deck.
    """

    db = SessionLocal()

    try:
        os.makedirs(GENERATED_DIR, exist_ok=True)
        os.makedirs(CHART_DIR, exist_ok=True)

        # --------------------------------------------------
        # SALES DATA
        # --------------------------------------------------

        finalized_bills = (
            db.query(Bill)
            .filter(Bill.status == "finalized")
            .all()
        )

        total_bills = len(finalized_bills)

        total_subtotal = sum(
            float(bill.subtotal or 0)
            for bill in finalized_bills
        )

        total_gst = sum(
            float(bill.gst_amount or 0)
            for bill in finalized_bills
        )

        total_sales = sum(
            float(bill.total_amount or 0)
            for bill in finalized_bills
        )

        average_bill = (
            total_sales / total_bills
            if total_bills > 0
            else 0
        )

        # --------------------------------------------------
        # PAYMENT ANALYSIS
        # --------------------------------------------------

        payment_data = {
            "cash": 0,
            "upi": 0,
            "card": 0,
            "credit": 0
        }

        payment_bill_counts = {
            "cash": 0,
            "upi": 0,
            "card": 0,
            "credit": 0
        }

        for bill in finalized_bills:

            mode = (bill.payment_mode or "cash").lower()

            amount = float(
                bill.total_amount or 0
            )

            if mode not in payment_data:
                payment_data[mode] = 0
                payment_bill_counts[mode] = 0

            payment_data[mode] += amount
            payment_bill_counts[mode] += 1

        # --------------------------------------------------
        # INVENTORY ANALYSIS
        # --------------------------------------------------

        products = (
            db.query(Product)
            .order_by(Product.quantity.asc())
            .all()
        )

        total_products = len(products)

        low_stock_products = [
            product
            for product in products
            if product.quantity <= product.reorder_level
        ]

        low_stock_count = len(low_stock_products)

        inventory_data = []

        for product in products:
            inventory_data.append({
                "name": product.name,
                "sku": product.sku,
                "quantity": product.quantity,
                "unit": product.unit,
                "reorder_level": product.reorder_level,
                "selling_price": product.selling_price
            })

        # --------------------------------------------------
        # CREDIT ANALYSIS
        # --------------------------------------------------

        credit_sales = payment_data.get(
            "credit",
            0
        )

        # --------------------------------------------------
        # CREATE CHARTS
        # --------------------------------------------------

        payment_chart = create_payment_chart(
            payment_data
        )

        inventory_chart = create_inventory_chart(
            inventory_data
        )

        # --------------------------------------------------
        # CREATE POWERPOINT
        # --------------------------------------------------

        presentation = Presentation()

        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        # --------------------------------------------------
        # SLIDE 1 - TITLE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_text_box(
            slide,
            STORE_NAME,
            0.8,
            2.2,
            11.7,
            1.0,
            32
        )

        add_text_box(
            slide,
            "Business Analysis & Operations Report",
            0.8,
            3.2,
            11.7,
            0.7,
            24
        )

        add_text_box(
            slide,
            datetime.now().strftime(
                "Generated on %d-%m-%Y %I:%M %p"
            ),
            0.8,
            4.2,
            11.7,
            0.5,
            14
        )

        # --------------------------------------------------
        # SLIDE 2 - SALES SUMMARY
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_title(
            slide,
            "Sales Summary",
            "Finalized bills currently stored in the supermarket database"
        )

        sales_text = (
            f"Total Bills: {total_bills}\n\n"
            f"Total Subtotal: {format_money(total_subtotal)}\n\n"
            f"Total GST: {format_money(total_gst)}\n\n"
            f"Total Sales: {format_money(total_sales)}\n\n"
            f"Average Bill Value: {format_money(average_bill)}"
        )

        add_text_box(
            slide,
            sales_text,
            1.0,
            1.7,
            11,
            4.5,
            22
        )

        # --------------------------------------------------
        # SLIDE 3 - PAYMENT ANALYSIS
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_title(
            slide,
            "Payment Mode Analysis"
        )

        slide.shapes.add_picture(
            payment_chart,
            Inches(0.8),
            Inches(1.4),
            width=Inches(7.0)
        )

        payment_text = "Payment Breakdown\n\n"

        for mode in payment_data:

            payment_text += (
                f"{mode.upper()}: "
                f"{format_money(payment_data[mode])} "
                f"({payment_bill_counts.get(mode, 0)} bill(s))\n\n"
            )

        add_text_box(
            slide,
            payment_text,
            8.1,
            1.6,
            4.5,
            4.5,
            16
        )

        # --------------------------------------------------
        # SLIDE 4 - INVENTORY
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_title(
            slide,
            "Inventory Analysis",
            "Products with the lowest stock levels are shown first"
        )

        slide.shapes.add_picture(
            inventory_chart,
            Inches(0.5),
            Inches(1.4),
            width=Inches(8.0)
        )

        inventory_text = (
            f"Total Product SKUs: {total_products}\n\n"
            f"Low Stock Products: {low_stock_count}\n\n"
        )

        if low_stock_products:

            inventory_text += "Reorder Candidates:\n\n"

            for product in low_stock_products[:5]:

                inventory_text += (
                    f"• {product.name}: "
                    f"{product.quantity} {product.unit}\n"
                )

        else:

            inventory_text += (
                "No products currently require reordering."
            )

        add_text_box(
            slide,
            inventory_text,
            8.7,
            1.6,
            4.0,
            4.8,
            15
        )

        # --------------------------------------------------
        # SLIDE 5 - PRODUCT STOCK TABLE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_title(
            slide,
            "Current Inventory Snapshot"
        )

        rows_to_show = products[:10]

        headers = [
            "Product",
            "SKU",
            "Stock",
            "Unit",
            "Reorder Level"
        ]

        rows = len(rows_to_show) + 1
        cols = len(headers)

        table = slide.shapes.add_table(
            rows,
            cols,
            Inches(0.5),
            Inches(1.4),
            Inches(12.3),
            Inches(5.2)
        ).table

        widths = [
            3.8,
            2.0,
            1.5,
            1.5,
            2.5
        ]

        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)

        for col_index, header in enumerate(headers):

            cell = table.cell(0, col_index)

            cell.text = header

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.bold = True
                paragraph.font.size = Pt(13)
                paragraph.alignment = PP_ALIGN.CENTER

        for row_index, product in enumerate(
            rows_to_show,
            start=1
        ):

            values = [
                product.name,
                product.sku,
                str(product.quantity),
                product.unit,
                str(product.reorder_level)
            ]

            for col_index, value in enumerate(values):

                cell = table.cell(
                    row_index,
                    col_index
                )

                cell.text = value

                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(11)
                    paragraph.alignment = PP_ALIGN.CENTER

        # --------------------------------------------------
        # SLIDE 6 - BUSINESS INSIGHTS
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        add_title(
            slide,
            "Business Insights"
        )

        insights = generate_business_insights(
            total_sales,
            total_bills,
            average_bill,
            credit_sales,
            low_stock_count,
            total_products
        )

        insight_text = ""

        for index, insight in enumerate(
            insights,
            start=1
        ):

            insight_text += (
                f"{index}. {insight}\n\n"
            )

        add_text_box(
            slide,
            insight_text,
            0.9,
            1.5,
            11.5,
            5.0,
            20
        )

        # --------------------------------------------------
        # SAVE PRESENTATION
        # --------------------------------------------------

        filename = "supermarket_business_analysis.pptx"

        filepath = os.path.join(
            GENERATED_DIR,
            filename
        )

        presentation.save(filepath)

        if not os.path.exists(filepath):

            return {
                "success": False,
                "message": (
                    "PPTX generation completed but "
                    "file was not found."
                )
            }

        file_size = os.path.getsize(filepath)

        if file_size == 0:

            return {
                "success": False,
                "message": "Generated PPTX file is empty."
            }

        return {
            "success": True,
            "message": (
                "Business analysis PPTX generated successfully."
            ),
            "filename": filename,
            "filepath": filepath,
            "file_size": file_size,
            "total_bills": total_bills,
            "total_sales": round(total_sales, 2),
            "total_gst": round(total_gst, 2),
            "average_bill": round(average_bill, 2),
            "low_stock_count": low_stock_count,
            "total_products": total_products
        }

    except Exception as e:

        return {
            "success": False,
            "message": (
                f"Error generating analysis deck: {str(e)}"
            )
        }

    finally:

        db.close()